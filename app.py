import streamlit as st
from pypdf import PdfReader
from docx import Document

from langchain_google_genai import (
    ChatGoogleGenerativeAI,
    GoogleGenerativeAIEmbeddings
)
from langchain_core.documents import Document as LCDocument
from langchain_community.vectorstores import FAISS
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain.tools import tool
from langchain.agents import create_agent


# --------------------------------------------------
# PAGE CONFIGURATION
# --------------------------------------------------

st.set_page_config(
    page_title="AI Resume Analyzer",
    page_icon="📄",
    layout="wide"
)


# --------------------------------------------------
# TITLE
# --------------------------------------------------

st.title("📄 AI Resume Analyzer Agent")

st.write(
    "Upload your resume and get an AI-powered analysis "
    "using Gemini, RAG and FAISS."
)


# --------------------------------------------------
# API KEY
# --------------------------------------------------

try:
    GOOGLE_API_KEY = st.secrets["GEMINI_API_KEY"]
except Exception:
    st.error("Gemini API key is not configured.")
    st.stop()


# --------------------------------------------------
# INITIALIZE GEMINI
# --------------------------------------------------

@st.cache_resource
def initialize_models():

    llm = ChatGoogleGenerativeAI(
        model="gemini-3.5-flash-lite",
        google_api_key=GOOGLE_API_KEY,
        temperature=0
    )

    embeddings = GoogleGenerativeAIEmbeddings(
        model="gemini-embedding-001",
        google_api_key=GOOGLE_API_KEY
    )

    return llm, embeddings


llm, embeddings = initialize_models()


# --------------------------------------------------
# EXTRACT RESUME TEXT
# --------------------------------------------------

def extract_resume_text(uploaded_file):

    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):

        reader = PdfReader(uploaded_file)

        text = ""

        for page in reader.pages:
            page_text = page.extract_text()

            if page_text:
                text += page_text + "\n"

        return text

    elif filename.endswith(".docx"):

        document = Document(uploaded_file)

        text = ""

        for paragraph in document.paragraphs:
            text += paragraph.text + "\n"

        return text

    else:

        raise ValueError(
            "Only PDF and DOCX files are supported."
        )


# --------------------------------------------------
# RESUME UPLOAD
# --------------------------------------------------

uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx"]
)


# --------------------------------------------------
# ANALYZE RESUME
# --------------------------------------------------

if uploaded_file is not None:

    st.success(
        f"Resume uploaded: {uploaded_file.name}"
    )

    if st.button("🔍 Analyze Resume"):

        with st.spinner("Analyzing your resume..."):

            try:

                # Extract text
                resume_text = extract_resume_text(
                    uploaded_file
                )

                if not resume_text.strip():
                    st.error(
                        "Could not extract text from the resume."
                    )
                    st.stop()

                # Split text
                text_splitter = RecursiveCharacterTextSplitter(
                    chunk_size=1000,
                    chunk_overlap=200
                )

                resume_chunks = text_splitter.split_text(
                    resume_text
                )

                # Create LangChain documents
                documents = [
                    LCDocument(
                        page_content=chunk,
                        metadata={
                            "source": uploaded_file.name
                        }
                    )
                    for chunk in resume_chunks
                ]

                # Create FAISS database
                vector_store = FAISS.from_documents(
                    documents,
                    embeddings
                )

                # Create retrieval tool
                @tool
                def retrieve_resume_context(query: str):
                    """
                    Retrieve relevant information
                    from the candidate's resume.
                    """

                    retrieved_docs = (
                        vector_store.similarity_search(
                            query,
                            k=3
                        )
                    )

                    serialized = "\n\n".join(
                        f"Content: {doc.page_content}"
                        for doc in retrieved_docs
                    )

                    return serialized

                # Agent instructions
                prompt = """
You are an AI Resume Analyzer Agent.

Your job is to analyze the candidate's resume
using information retrieved from the resume.

Provide the following sections:

1. Technical Skills
2. Education
3. Projects
4. Work Experience
5. Certifications
6. Achievements
7. Overall Resume Summary

Important rules:

- Use only information present in the resume.
- Do not invent information.
- Keep the analysis clear and structured.
- If a section is not present in the resume, say:
  "Not mentioned in the resume."
"""

                # Create agent
                tools = [
                    retrieve_resume_context
                ]

                resume_agent = create_agent(
                    llm,
                    tools,
                    system_prompt=prompt
                )

                # Ask agent to analyze
                response = resume_agent.invoke({
                    "messages": [
                        {
                            "role": "user",
                            "content": """
Analyze this candidate's resume.

Provide:

1. Technical Skills
2. Education
3. Projects
4. Work Experience
5. Certifications
6. Achievements
7. Overall Resume Summary

Use only information available in the resume.
Do not invent any information.
"""
                        }
                    ]
                })

                # Display result
                analysis = (
                    response["messages"][-1].content
                )

                st.success(
                    "Resume analysis completed!"
                )

                st.markdown("## 📊 Resume Analysis")

                st.markdown(analysis)

            except Exception as e:

                st.error(
                    f"An error occurred: {str(e)}"
                )


# --------------------------------------------------
# FOOTER
# --------------------------------------------------

st.markdown("---")

st.caption(
    "AI Resume Analyzer Agent | Gemini + LangChain + RAG + FAISS"
)
